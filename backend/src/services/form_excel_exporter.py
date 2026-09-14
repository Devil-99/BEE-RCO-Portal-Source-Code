from openpyxl import load_workbook
from fastapi.responses import StreamingResponse
import io
from src.models import SubmissionDetails

TEMPLATE_PATH = "src/templates/cpp_submission_template.xlsx"

def read_excel_mapping(wb):
    meta = wb["_META_"]

    mapping = []
    for row in meta.iter_rows(min_row=2, values_only=True):
        acronym, sheet, cell = row
        mapping.append({
            "acronym": acronym,
            "sheet": sheet,
            "cell": cell
        })
    return mapping

def fill_excel_from_submission(template_path, submission_data):
    wb = load_workbook(template_path)
    mapping = read_excel_mapping(wb)

    for m in mapping:
        acronym = m["acronym"]
        if acronym in submission_data:
            ws = wb[m["sheet"]]
            ws[m["cell"]] = submission_data[acronym]

    return wb

def get_submission_data(session, submission_id):
    rows = (
        session.query(SubmissionDetails.acronym, SubmissionDetails.value)
        .filter(SubmissionDetails.submission_id == submission_id)
        .all()
    )

    return {r.acronym: float(r.value) for r in rows}


def export_submission_excel(session, submission_id):
    data = get_submission_data(session, submission_id)

    wb = fill_excel_from_submission(TEMPLATE_PATH, data)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer
