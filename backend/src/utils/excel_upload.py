import io
import logging
from datetime import date, datetime
from typing import Optional

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook


DEFAULT_MAX_UPLOAD_ROWS = 5000


def clean_string(value):
    if value is None:
        return None

    value = str(value).strip()

    return value if value else None


def parse_excel_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):
        try:
            return datetime.strptime(
                value.strip(),
                "%Y-%m-%d"
            ).date()

        except ValueError:
            return None

    return None


def read_excel_upload(
    file: UploadFile,
    template_headers: list[str],
    max_upload_rows: int = DEFAULT_MAX_UPLOAD_ROWS,
    logger: Optional[logging.Logger] = None,
    log_prefix: str = "",
):
    # Validate file extension

    if (
        not file.filename
        or not file.filename.lower().endswith(".xlsx")
    ):
        if logger:
            logger.warning(
                f"{log_prefix}Invalid file format - only .xlsx allowed"
            )

        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are allowed."
        )

    # Read uploaded file

    contents = file.file.read()

    if not contents:
        if logger:
            logger.warning(
                f"{log_prefix}Empty file uploaded"
            )

        raise HTTPException(
            status_code=400,
            detail="Empty file. Please upload a valid Excel file."
        )

    # Load workbook

    try:
        workbook = load_workbook(
            io.BytesIO(contents)
        )

        worksheet = workbook.active

    except Exception:
        if logger:
            logger.exception(
                f"{log_prefix}Failed to read uploaded excel file"
            )

        raise HTTPException(
            status_code=400,
            detail="Invalid Excel file format."
        )

    # Validate headers

    uploaded_headers = [
        cell.value
        for cell in worksheet[1]
    ]

    if uploaded_headers != template_headers:
        if logger:
            logger.warning(
                f"{log_prefix}Invalid template headers"
            )

        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid template. "
                "Please download the latest template and try again."
            )
        )

    # Read rows

    rows = list(
        worksheet.iter_rows(
            min_row=2,
            values_only=True
        )
    )

    # Remove completely empty rows

    non_empty_rows = [
        row
        for row in rows
        if any(
            cell is not None
            and str(cell).strip()
            for cell in row
        )
    ]

    # Validate records exist

    if not non_empty_rows:
        if logger:
            logger.warning(
                f"{log_prefix}No records found in uploaded file"
            )

        raise HTTPException(
            status_code=400,
            detail="No records found in uploaded file."
        )

    # Validate maximum rows

    if len(non_empty_rows) > max_upload_rows:
        if logger:
            logger.warning(
                f"{log_prefix}Maximum {max_upload_rows} "
                f"records allowed per upload"
            )

        raise HTTPException(
            status_code=400,
            detail=(
                f"Maximum {max_upload_rows} "
                f"records allowed per upload."
            )
        )

    return non_empty_rows


def map_excel_row(
    headers: list[str],
    header_to_field: dict[str, str],
    excel_row: tuple,
):

    return {
        header_to_field[header]: value
        for header, value in zip(
            headers,
            excel_row
        )
    }


def raise_validation_errors(
    validation_errors: list[str],
):


    if not validation_errors:
        return

    raise HTTPException(
        status_code=400,
        detail={
            "message": "Validation Failed",
            "errors": validation_errors,
        }
    )